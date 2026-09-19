import streamlit as st
from io import BytesIO


# ==========================================================
# ATS RESUME TEMPLATES
# ==========================================================

ATS_TEMPLATES = {
    "Fresher / Student": {
        "description": (
            "One-page ATS-oriented template for students and fresh graduates "
            "with education, skills, projects, certifications and achievements."
        ),
        "sections": [
            "Contact Information",
            "Professional Summary",
            "Education",
            "Technical Skills",
            "Projects",
            "Certifications",
            "Achievements",
        ],
        "content": """YOUR NAME
City, State | +91 XXXXX XXXXX | email@example.com
LinkedIn: linkedin.com/in/yourprofile | GitHub: github.com/yourusername

PROFESSIONAL SUMMARY

Computer Science / AI & ML student with strong foundations in software
development, Data Structures & Algorithms, databases, and Artificial
Intelligence. Experienced in building practical projects using modern
programming languages and development frameworks.

EDUCATION

Bachelor of Technology (B.Tech) – Computer Science & Engineering (AI & ML)
University / Institute Name | City, State
20XX – Present

TECHNICAL SKILLS

Programming Languages: Java, Python, JavaScript, SQL
Web Development: HTML, CSS, React.js, Node.js, Express.js
Backend: REST APIs, FastAPI
Databases: MongoDB, MySQL
AI / ML: Machine Learning, Artificial Intelligence, NLP
Core CS: Data Structures & Algorithms, OOP, DBMS
Tools: Git, GitHub, VS Code, Google Colab

PROJECTS

PROJECT NAME | Technology 1, Technology 2, Technology 3

• Developed and implemented [project functionality].
• Designed [feature/system] to solve [specific problem].
• Integrated [API/database/model] into the application.
• Improved [process/result] through [technical approach].

PROJECT NAME | Technology 1, Technology 2

• Built [application/system] for [purpose].
• Implemented [important feature].
• Used [technology] to handle [specific requirement].

CERTIFICATIONS

• Certification Name – Issuing Organization
• Certification Name – Issuing Organization
• Certification Name – Issuing Organization

ACHIEVEMENTS

• Research publication / technical achievement.
• Patent / design registration / innovation contribution.
• Hackathon / coding / academic achievement.
""",
    },

    "Software Engineer": {
        "description": (
            "General-purpose ATS-oriented template for Software Engineer, "
            "SDE and Trainee Software Engineer applications."
        ),
        "sections": [
            "Contact Information",
            "Professional Summary",
            "Technical Skills",
            "Experience",
            "Projects",
            "Education",
            "Certifications",
        ],
        "content": """YOUR NAME
City, State | +91 XXXXX XXXXX | email@example.com
LinkedIn: linkedin.com/in/yourprofile | GitHub: github.com/yourusername

PROFESSIONAL SUMMARY

Aspiring Software Engineer with strong foundations in Java, Python,
Data Structures & Algorithms, Object-Oriented Programming, databases,
and full-stack development. Experienced in designing and implementing
software projects using modern development tools and APIs.

TECHNICAL SKILLS

Languages: Java, Python, JavaScript, SQL
Frontend: HTML, CSS, JavaScript, React.js
Backend: Node.js, Express.js, FastAPI, REST APIs
Databases: MongoDB, MySQL, PostgreSQL
Core CS: Data Structures & Algorithms, OOP, DBMS, Operating Systems
AI / ML: Machine Learning, NLP, Artificial Intelligence
Tools: Git, GitHub, VS Code, Postman

EXPERIENCE

JOB TITLE | COMPANY NAME
City, State | Month Year – Month Year

• Developed [application/feature] using [technology].
• Implemented [technical functionality] to address [problem].
• Collaborated with team members to deliver [project/result].
• Debugged and improved [system/process] using [approach].

PROJECTS

PROJECT NAME | Java, Spring Boot, SQL

• Developed [application] for [purpose].
• Implemented RESTful APIs for [functionality].
• Designed database structures for [data requirement].
• Applied OOP and DSA concepts during implementation.

PROJECT NAME | Python, FastAPI, Machine Learning

• Developed an AI-powered system for [purpose].
• Implemented data processing and model integration.
• Created API endpoints for application functionality.

EDUCATION

Bachelor of Technology (B.Tech) – Computer Science & Engineering
University / Institute Name | City, State
20XX – 20XX

CERTIFICATIONS

• Certification Name – Issuing Organization
• Certification Name – Issuing Organization
""",
    },

    "AI / ML Engineer": {
        "description": (
            "ATS-oriented template focused on Artificial Intelligence, "
            "Machine Learning, NLP, Python and AI-powered projects."
        ),
        "sections": [
            "Contact Information",
            "Professional Summary",
            "Technical Skills",
            "AI / ML Projects",
            "Research",
            "Education",
            "Certifications",
        ],
        "content": """YOUR NAME
City, State | +91 XXXXX XXXXX | email@example.com
LinkedIn: linkedin.com/in/yourprofile | GitHub: github.com/yourusername

PROFESSIONAL SUMMARY

Aspiring AI/ML Engineer with practical knowledge of Python, Machine
Learning, Artificial Intelligence, Natural Language Processing and
software development. Experienced in developing AI-powered applications,
integrating models into software systems, and working with data-processing
and backend technologies.

TECHNICAL SKILLS

Programming: Python, Java, JavaScript, SQL
Machine Learning: Machine Learning, Supervised Learning, Unsupervised Learning
AI / NLP: Artificial Intelligence, NLP, Natural Language Processing
Libraries: NumPy, Pandas, Scikit-learn
Deep Learning: TensorFlow / PyTorch
Backend: FastAPI, Flask, Streamlit
Databases: SQL, MongoDB
Tools: Git, GitHub, Google Colab, JupyterLab, VS Code
Core CS: DSA, OOP, DBMS

AI / ML PROJECTS

AI PROJECT NAME | Python, Machine Learning, NLP

• Developed an AI-based system for [problem].
• Performed data preprocessing and feature preparation.
• Applied [algorithm/model] for [task].
• Evaluated model performance using [metric].
• Integrated the model into an application using [framework].

AI PROJECT NAME | Python, Computer Vision

• Developed a computer vision application for [purpose].
• Implemented image preprocessing and recognition functionality.
• Integrated the trained model into the application workflow.

RESEARCH

RESEARCH PAPER TITLE
Publication / Journal / Conference | Year

• Researched [topic/problem].
• Proposed and evaluated [approach/system].
• Documented findings and implementation results.

EDUCATION

Bachelor of Technology (B.Tech) – Computer Science & Engineering (AI & ML)
University / Institute Name | City, State
20XX – 20XX

CERTIFICATIONS

• Introduction to Machine Learning – IIT / Organization
• Artificial Intelligence Foundations
• Machine Learning / AI Certification
""",
    },

    "Full Stack Developer": {
        "description": (
            "ATS-oriented template emphasizing frontend, backend, APIs, "
            "databases, authentication and full-stack projects."
        ),
        "sections": [
            "Contact Information",
            "Professional Summary",
            "Technical Skills",
            "Experience",
            "Projects",
            "Education",
            "Certifications",
        ],
        "content": """YOUR NAME
City, State | +91 XXXXX XXXXX | email@example.com
LinkedIn: linkedin.com/in/yourprofile | GitHub: github.com/yourusername

PROFESSIONAL SUMMARY

Aspiring Full Stack Developer with experience building responsive web
applications, RESTful APIs, backend services and database-driven systems.
Skilled in JavaScript, React.js, Node.js, Express.js, Python, FastAPI,
SQL and MongoDB.

TECHNICAL SKILLS

Languages: Java, JavaScript, Python, SQL
Frontend: HTML5, CSS3, JavaScript, React.js, Bootstrap
Backend: Node.js, Express.js, FastAPI
APIs: REST API, RESTful Services, API Integration
Databases: MongoDB, MySQL, PostgreSQL
Authentication: Sessions, JWT, OAuth
Tools: Git, GitHub, VS Code, Postman
Core CS: DSA, OOP, DBMS

EXPERIENCE

JOB TITLE | COMPANY NAME
City, State | Month Year – Month Year

• Developed frontend and backend functionality using [technology].
• Created REST APIs for [application requirement].
• Integrated [database/API/service] into the application.
• Debugged and improved application functionality.

PROJECTS

FULL STACK PROJECT | React.js, Node.js, Express.js, MongoDB

• Developed a responsive full-stack web application for [purpose].
• Created RESTful APIs for [functionality].
• Implemented MongoDB data storage and retrieval.
• Built reusable frontend components using React.js.
• Used Git and GitHub for version control.

WEB APPLICATION | HTML, CSS, JavaScript, Node.js

• Developed [application] to solve [problem].
• Implemented [feature].
• Integrated backend services with frontend functionality.

EDUCATION

Bachelor of Technology (B.Tech) – Computer Science & Engineering
University / Institute Name | City, State
20XX – 20XX

CERTIFICATIONS

• Full Stack Web Development
• Data Structures & Algorithms
• Database / SQL Certification
""",
    },
}


# ==========================================================
# DOCX GENERATOR
# ==========================================================

def _build_docx(title: str, content: str) -> bytes | None:
    """
    Generate a simple single-column ATS-oriented DOCX.
    """

    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.section import WD_SECTION
    except ImportError:
        return None

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    # Normal text
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10.5)

    # Main title
    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title_paragraph.add_run(
        title.upper()
    )

    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(15)

    # Resume content
    for raw_line in content.splitlines():

        line = raw_line.strip()

        if not line:
            document.add_paragraph()
            continue

        # Standard section heading
        if (
            line.isupper()
            and len(line) <= 50
            and not line.startswith("YOUR NAME")
            and "|" not in line
        ):
            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(2)

            run = paragraph.add_run(line)
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(11)

        else:
            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(2)

            run = paragraph.add_run(line)
            run.font.name = "Arial"
            run.font.size = Pt(10.5)

    output = BytesIO()
    document.save(output)

    return output.getvalue()


# ==========================================================
# BULLET GENERATOR
# ==========================================================

def generate_bullet(role, action, technology, result):

    action = action.strip()
    technology = technology.strip()
    result = result.strip()

    if not action:
        return ""

    if technology and result:
        return (
            f"Developed and implemented {action} using "
            f"{technology}, resulting in {result}."
        )

    if technology:
        return (
            f"Developed and implemented {action} using "
            f"{technology}."
        )

    if result:
        return (
            f"Developed and implemented {action}, "
            f"resulting in {result}."
        )

    return f"Developed and implemented {action}."


# ==========================================================
# PROFESSIONAL SUMMARY
# ==========================================================

def generate_summary(role, experience, skills, projects):

    skills_text = ", ".join(skills)

    if experience == "Fresher / Student":
        opening = (
            f"Aspiring {role} with strong foundations in "
            f"{skills_text}."
        )
    else:
        opening = (
            f"{experience} {role} with experience in "
            f"{skills_text}."
        )

    return (
        f"{opening} Experienced in developing software projects "
        f"and applying problem-solving skills to real-world "
        f"applications. Built {projects} and demonstrated "
        f"practical knowledge through hands-on development."
    )


# ==========================================================
# JOB DESCRIPTION SKILL DATABASE
# ==========================================================

JOB_SKILL_DATABASE = {

    "Programming Languages": [
        "Python",
        "Java",
        "JavaScript",
        "TypeScript",
        "C++",
        "C#",
        "Go",
        "Rust",
        "SQL",
    ],

    "Frontend": [
        "HTML",
        "CSS",
        "React",
        "React.js",
        "Angular",
        "Vue",
        "Next.js",
        "Bootstrap",
        "Tailwind",
    ],

    "Backend": [
        "Node.js",
        "Express.js",
        "Django",
        "Flask",
        "FastAPI",
        "Spring",
        "Spring Boot",
        "REST API",
        "RESTful API",
    ],

    "Databases": [
        "MySQL",
        "PostgreSQL",
        "MongoDB",
        "Redis",
        "SQLite",
        "Oracle",
        "SQL Server",
    ],

    "AI / ML": [
        "Machine Learning",
        "Artificial Intelligence",
        "Deep Learning",
        "Natural Language Processing",
        "NLP",
        "TensorFlow",
        "PyTorch",
        "Scikit-learn",
        "Pandas",
        "NumPy",
        "Computer Vision",
        "Generative AI",
        "LLM",
    ],

    "Cloud / DevOps": [
        "AWS",
        "Azure",
        "Google Cloud",
        "Docker",
        "Kubernetes",
        "CI/CD",
        "Jenkins",
        "GitHub Actions",
    ],

    "Core CS": [
        "Data Structures",
        "Algorithms",
        "Data Structures and Algorithms",
        "Object Oriented Programming",
        "OOP",
        "DBMS",
        "Operating Systems",
        "Computer Networks",
        "System Design",
    ],

    "Tools": [
        "Git",
        "GitHub",
        "VS Code",
        "Jira",
        "Postman",
        "Figma",
    ],

    "Soft Skills": [
        "Communication",
        "Leadership",
        "Teamwork",
        "Problem Solving",
        "Time Management",
        "Collaboration",
        "Adaptability",
    ],
}


# ==========================================================
# JOB DESCRIPTION ANALYZER
# ==========================================================

def extract_jd_skills(job_description):

    if not job_description:
        return {}

    jd_lower = job_description.lower()

    detected = {}

    for category, skills in JOB_SKILL_DATABASE.items():

        found = []

        for skill in skills:

            if skill.lower() in jd_lower:
                found.append(skill)

        if found:
            detected[category] = found

    return detected


# ==========================================================
# SKILL GAP CALCULATOR
# ==========================================================

def calculate_skill_match(required_skills, user_skills):

    required_normalized = {
        skill.lower().strip()
        for skill in required_skills
    }

    user_normalized = {
        skill.lower().strip()
        for skill in user_skills
    }

    matched = []
    missing = []

    for skill in required_skills:

        if skill.lower().strip() in user_normalized:
            matched.append(skill)
        else:
            missing.append(skill)

    if required_normalized:

        percentage = round(
            len(
                required_normalized.intersection(
                    user_normalized
                )
            )
            / len(required_normalized)
            * 100
        )

    else:
        percentage = 0

    return matched, missing, percentage



# ==========================================================
# JOB DESCRIPTION ANALYZER
# ==========================================================

JOB_SKILL_DATABASE = {

    "Programming Languages": [
        "Python",
        "Java",
        "JavaScript",
        "TypeScript",
        "C++",
        "C#",
        "Go",
        "Rust",
        "SQL",
    ],

    "Frontend": [
        "HTML",
        "CSS",
        "React",
        "React.js",
        "Angular",
        "Vue",
        "Next.js",
        "Bootstrap",
        "Tailwind",
    ],

    "Backend": [
        "Node.js",
        "Express.js",
        "Django",
        "Flask",
        "FastAPI",
        "Spring",
        "Spring Boot",
        "REST API",
        "RESTful API",
    ],

    "Databases": [
        "MySQL",
        "PostgreSQL",
        "MongoDB",
        "Redis",
        "SQLite",
        "Oracle",
        "SQL Server",
    ],

    "AI / ML": [
        "Machine Learning",
        "Artificial Intelligence",
        "Deep Learning",
        "Natural Language Processing",
        "NLP",
        "TensorFlow",
        "PyTorch",
        "Scikit-learn",
        "Pandas",
        "NumPy",
        "Computer Vision",
        "Generative AI",
        "LLM",
    ],

    "Cloud / DevOps": [
        "AWS",
        "Azure",
        "Google Cloud",
        "Docker",
        "Kubernetes",
        "CI/CD",
        "Jenkins",
        "GitHub Actions",
    ],

    "Core CS": [
        "Data Structures",
        "Algorithms",
        "Data Structures and Algorithms",
        "Object Oriented Programming",
        "OOP",
        "DBMS",
        "Operating Systems",
        "Computer Networks",
        "System Design",
    ],

    "Tools": [
        "Git",
        "GitHub",
        "VS Code",
        "Jira",
        "Postman",
        "Figma",
    ],

    "Soft Skills": [
        "Communication",
        "Leadership",
        "Teamwork",
        "Problem Solving",
        "Time Management",
        "Collaboration",
        "Adaptability",
    ],
}


def extract_jd_skills(job_description):

    if not job_description:
        return {}

    jd_lower = job_description.lower()

    detected = {}

    for category, skills in JOB_SKILL_DATABASE.items():

        found = []

        for skill in skills:

            # Case-insensitive search
            if skill.lower() in jd_lower:
                found.append(skill)

        if found:
            detected[category] = found

    return detected


def calculate_skill_match(required_skills, user_skills):

    required_normalized = {
        skill.lower().strip()
        for skill in required_skills
    }

    user_normalized = {
        skill.lower().strip()
        for skill in user_skills
    }

    matched = []
    missing = []

    for skill in required_skills:

        if skill.lower().strip() in user_normalized:
            matched.append(skill)
        else:
            missing.append(skill)

    if required_normalized:
        percentage = round(
            len(required_normalized.intersection(user_normalized))
            / len(required_normalized)
            * 100
        )
    else:
        percentage = 0

    return matched, missing, percentage


# ==========================================================
# ATS KEYWORDS
# ==========================================================

ATS_KEYWORDS = {

    "Software Engineer": [
        "Java",
        "Python",
        "JavaScript",
        "Data Structures",
        "Algorithms",
        "OOP",
        "REST API",
        "Git",
        "GitHub",
        "SQL",
        "MongoDB",
        "Testing",
        "Debugging",
        "Agile",
        "CI/CD",
    ],

    "Full Stack Developer": [
        "React",
        "JavaScript",
        "HTML",
        "CSS",
        "Node.js",
        "Express.js",
        "REST API",
        "MongoDB",
        "SQL",
        "Git",
        "Authentication",
        "API Integration",
        "Responsive Design",
        "Deployment",
    ],

    "AI / ML Engineer": [
        "Python",
        "Machine Learning",
        "Artificial Intelligence",
        "NLP",
        "Deep Learning",
        "Scikit-learn",
        "Pandas",
        "NumPy",
        "Model Evaluation",
        "Data Preprocessing",
        "Feature Engineering",
        "FastAPI",
        "Streamlit",
        "TensorFlow",
    ],

    "Data Analyst": [
        "Python",
        "SQL",
        "Excel",
        "Power BI",
        "Tableau",
        "Data Analysis",
        "Data Visualization",
        "Statistics",
        "Pandas",
        "NumPy",
        "Reporting",
        "Dashboard",
    ],
}


# ==========================================================
# ATS CHECKLIST
# ==========================================================

def ats_checklist():

    st.markdown("### ✅ ATS Resume Submission Checklist")

    checklist = [
        "Resume uses a single-column layout",
        "Standard headings are used",
        "No tables or text boxes",
        "No unnecessary graphics or icons",
        "Contact information is clearly visible",
        "Target job keywords are included",
        "Technical skills match the target role",
        "Projects contain measurable contributions",
        "Action verbs are used in bullet points",
        "Dates are consistently formatted",
        "Resume contains no spelling errors",
        "PDF has been checked after export",
    ]

    completed = 0

    for index, item in enumerate(checklist):

        checked = st.checkbox(
            item,
            key=f"ats_check_{index}",
        )

        if checked:
            completed += 1

    percentage = int(
        (completed / len(checklist)) * 100
    )

    st.progress(percentage / 100)

    st.markdown(
        f"### ATS Readiness: {percentage}%"
    )

    if percentage == 100:

        st.success(
            "🎉 Checklist complete! "
            "Your resume is ready for final review."
        )

    elif percentage >= 75:

        st.info(
            "Almost there! Complete the remaining "
            "checks before submitting."
        )

    else:

        st.warning(
            "Complete more checks before submitting "
            "your resume."
        )


# ==========================================================
# ATS RESUME TEMPLATE TAB
# ==========================================================

def ats_resume_templates():

    st.markdown("### 📄 ATS-Friendly Resume Templates")

    st.write(
        "Choose a clean, single-column resume template "
        "and download an editable DOCX or plain-text version."
    )

    st.info(
        "ATS-oriented design: single column, standard section "
        "headings, no tables, no sidebars, no graphics, and "
        "contact information in the document body."
    )

    selected_template = st.selectbox(
        "Choose a resume template",
        list(ATS_TEMPLATES.keys()),
        key="ats_template_selector",
    )

    template = ATS_TEMPLATES[selected_template]

    st.markdown(
        f"## 📄 {selected_template}"
    )

    st.write(template["description"])

    st.markdown("#### ATS Structure")

    structure_columns = st.columns(4)

    structure_items = [
        "Single column",
        "Standard headings",
        "No tables",
        "Text-based content",
    ]

    for index, item in enumerate(structure_items):

        with structure_columns[index]:

            st.success(
                f"✓ {item}"
            )

    st.divider()

    # ------------------------------------------------------
    # Preview
    # ------------------------------------------------------

    st.markdown("### 👁️ Template Preview")

    st.code(
        template["content"],
        language="text",
    )

    st.divider()

    # ------------------------------------------------------
    # Sections
    # ------------------------------------------------------

    st.markdown("### 📑 Recommended Sections")

    section_columns = st.columns(2)

    for index, section_name in enumerate(
        template["sections"]
    ):

        with section_columns[index % 2]:

            st.markdown(
                f"✓ {section_name}"
            )

    st.divider()

    # ------------------------------------------------------
    # Downloads
    # ------------------------------------------------------

    st.markdown("### ⬇️ Download Template")

    txt_bytes = template["content"].encode(
        "utf-8"
    )

    docx_bytes = _build_docx(
        selected_template,
        template["content"],
    )

    safe_name = (
        selected_template
        .lower()
        .replace("/", "-")
        .replace(" ", "_")
        .replace("&", "and")
    )

    download_columns = st.columns(2)

    with download_columns[0]:

        st.download_button(
            label="⬇️ Download DOCX",
            data=(
                docx_bytes
                if docx_bytes is not None
                else txt_bytes
            ),
            file_name=(
                f"{safe_name}_ATS_Resume_Template."
                f"{'docx' if docx_bytes is not None else 'txt'}"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
                if docx_bytes is not None
                else "text/plain"
            ),
            use_container_width=True,
        )

    with download_columns[1]:

        st.download_button(
            label="⬇️ Download TXT",
            data=txt_bytes,
            file_name=(
                f"{safe_name}_ATS_Resume_Template.txt"
            ),
            mime="text/plain",
            use_container_width=True,
        )

    st.divider()

    st.markdown("### ⚠️ Before Using the Template")

    st.warning(
        "Replace all placeholder text with your own information. "
        "Only include skills, technologies, experience and "
        "achievements that are genuinely applicable to you."
    )

    st.caption(
        "ATS compatibility varies across employers and ATS "
        "implementations. Always review the final exported "
        "document before submitting it."
    )


# ==========================================================
# MAIN PAGE
# ==========================================================

def show():

    st.markdown(
        """
        <div class="page-header">
            <h1>📚 Career & ATS Toolkit</h1>
            <p>
                Build stronger resumes, analyze job descriptions,
                discover ATS keywords, and prepare your application
                before submission.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    # ======================================================
    # QUICK CAREER TOOLS
    # ======================================================

    st.markdown("## 🛠️ Quick Career Tools")

    tool_tabs = st.tabs(
        [
            "📝 Bullet Generator",
            "✨ Summary Builder",
            "🎯 JD Analyzer",
            "🧩 Skill Gap",
            "🔑 Keyword Explorer",
            "✅ ATS Checklist",
            "📄 ATS Resume Templates",
        ]
    )

    # ======================================================
    # BULLET GENERATOR
    # ======================================================

    with tool_tabs[0]:

        st.markdown(
            "### 📝 Resume Bullet Generator"
        )

        st.write(
            "Turn a simple project description into "
            "a more professional resume bullet."
        )

        action = st.text_input(
            "What did you build or do?",
            placeholder=(
                "e.g. an AI resume screening system"
            ),
            key="bullet_action",
        )

        technology = st.text_input(
            "Technologies used",
            placeholder=(
                "e.g. Python, FastAPI, Streamlit, NLP"
            ),
            key="bullet_technology",
        )

        result = st.text_input(
            "Result / impact",
            placeholder=(
                "e.g. automated resume analysis"
            ),
            key="bullet_result",
        )

        if st.button(
            "✨ Generate Resume Bullet",
            use_container_width=True,
            key="generate_bullet_button",
        ):

            if not action:

                st.warning(
                    "Please describe what you built or did."
                )

            else:

                bullet = generate_bullet(
                    "",
                    action,
                    technology,
                    result,
                )

                st.success(
                    "Generated Resume Bullet"
                )

                st.markdown(
                    f"> **{bullet}**"
                )

                st.code(bullet)


    # ======================================================
    # SUMMARY BUILDER
    # ======================================================

    with tool_tabs[1]:

        st.markdown(
            "### ✨ Professional Summary Builder"
        )

        role = st.selectbox(
            "Target Role",
            [
                "Software Engineer",
                "Full Stack Developer",
                "AI/ML Engineer",
                "Data Analyst",
                "Backend Developer",
                "Frontend Developer",
            ],
            key="summary_role",
        )

        experience = st.selectbox(
            "Experience Level",
            [
                "Fresher / Student",
                "1+ year experience",
                "2+ years experience",
                "3+ years experience",
            ],
            key="summary_experience",
        )

        skills = st.multiselect(
            "Select your strongest skills",
            [
                "Java",
                "Python",
                "JavaScript",
                "React.js",
                "Node.js",
                "FastAPI",
                "SQL",
                "MongoDB",
                "Machine Learning",
                "Artificial Intelligence",
                "NLP",
                "Data Structures & Algorithms",
                "Git",
            ],
            key="summary_skills",
        )

        projects = st.text_input(
            "Number/type of projects",
            placeholder=(
                "e.g. 3 AI and full-stack projects"
            ),
            key="summary_projects",
        )

        if st.button(
            "✨ Generate Professional Summary",
            use_container_width=True,
            key="generate_summary_button",
        ):

            if not skills:

                st.warning(
                    "Select at least one skill."
                )

            elif not projects:

                st.warning(
                    "Enter your project experience."
                )

            else:

                summary = generate_summary(
                    role,
                    experience,
                    skills,
                    projects,
                )

                st.success(
                    "Generated Summary"
                )

                st.markdown(
                    f"> {summary}"
                )

                st.code(summary)


    # ======================================================
    # JOB DESCRIPTION ANALYZER
    # ======================================================

    with tool_tabs[2]:

        st.markdown(
            "### 🎯 Job Description Analyzer"
        )

        st.write(
            "Paste a job description and identify "
            "important technical and professional requirements."
        )

        job_description = st.text_area(
            "Paste Job Description",
            height=280,
            placeholder=(
                "Example:\n\n"
                "We are looking for a Software Engineer "
                "with experience in Java, Python, React, "
                "REST APIs, SQL, Git and Docker..."
            ),
            key="jd_analyzer_text",
        )

        if st.button(
            "🔍 Analyze Job Description",
            use_container_width=True,
            key="analyze_jd_button",
        ):

            if not job_description.strip():

                st.warning(
                    "Please paste a job description first."
                )

            else:

                detected = extract_jd_skills(
                    job_description
                )

                total_skills = sum(
                    len(skills)
                    for skills in detected.values()
                )

                st.success(
                    f"Detected {total_skills} relevant skills."
                )

                if not detected:

                    st.warning(
                        "No predefined skills were detected. "
                        "Try a more detailed job description."
                    )

                else:

                    st.markdown(
                        "#### 🔑 Detected Job Requirements"
                    )

                    for category, detected_skills in (
                        detected.items()
                    ):

                        st.markdown(
                            f"**{category}**"
                        )

                        columns = st.columns(3)

                        for index, skill in enumerate(
                            detected_skills
                        ):

                            with columns[index % 3]:

                                st.success(
                                    f"✓ {skill}"
                                )

                    st.session_state[
                        "jd_detected_skills"
                    ] = [
                        skill
                        for detected_skills in detected.values()
                        for skill in detected_skills
                    ]

                    st.info(
                        "Switch to the 🧩 Skill Gap tab "
                        "to compare these requirements "
                        "with your skills."
                    )


    # ======================================================
    # SKILL GAP ANALYZER
    # ======================================================

    with tool_tabs[3]:

        st.markdown(
            "### 🧩 Skill Gap Analyzer"
        )

        st.write(
            "Compare skills detected from a job description "
            "with your current skills."
        )

        jd_skills = st.session_state.get(
            "jd_detected_skills",
            [],
        )

        if not jd_skills:

            st.info(
                "First analyze a Job Description in "
                "the 🎯 JD Analyzer tab."
            )

        else:

            st.markdown(
                f"**{len(jd_skills)} skills detected from the JD.**"
            )

            user_skill_options = sorted(
                {
                    skill
                    for skills in JOB_SKILL_DATABASE.values()
                    for skill in skills
                }
            )

            user_skills = st.multiselect(
                "Select the skills you currently have",
                user_skill_options,
                key="user_skill_gap_selection",
            )

            if st.button(
                "📊 Calculate Skill Gap",
                use_container_width=True,
                key="calculate_skill_gap",
            ):

                matched, missing, percentage = (
                    calculate_skill_match(
                        jd_skills,
                        user_skills,
                    )
                )

                st.markdown(
                    "### 📊 Skill Match"
                )

                col1, col2, col3 = st.columns(3)

                with col1:

                    st.metric(
                        "JD Skills",
                        len(jd_skills),
                    )

                with col2:

                    st.metric(
                        "Skills Matched",
                        len(matched),
                    )

                with col3:

                    st.metric(
                        "Skills Missing",
                        len(missing),
                    )

                st.progress(
                    percentage / 100
                )

                st.markdown(
                    f"### 🎯 Skill Match: {percentage}%"
                )

                if matched:

                    st.markdown(
                        "#### ✅ Skills You Have"
                    )

                    for skill in matched:

                        st.success(
                            f"✓ {skill}"
                        )

                if missing:

                    st.markdown(
                        "#### ⚠️ Skills Missing"
                    )

                    for skill in missing:

                        st.warning(
                            f"⚠️ {skill}"
                        )

                    st.info(
                        "Only add a skill to your resume "
                        "if you genuinely have that skill."
                    )

                else:

                    st.success(
                        "🎉 All detected JD skills are "
                        "present in your selected skill set."
                    )


    # ======================================================
    # KEYWORD EXPLORER
    # ======================================================

    with tool_tabs[4]:

        st.markdown(
            "### 🔑 ATS Keyword Explorer"
        )

        selected_role = st.selectbox(
            "Choose your target role",
            list(ATS_KEYWORDS.keys()),
            key="keyword_role",
        )

        keywords = ATS_KEYWORDS[
            selected_role
        ]

        st.write(
            f"Important ATS keywords commonly associated "
            f"with **{selected_role}**:"
        )

        columns = st.columns(3)

        for index, keyword in enumerate(
            keywords
        ):

            with columns[index % 3]:

                st.checkbox(
                    keyword,
                    value=False,
                    key=(
                        f"keyword_"
                        f"{selected_role}_"
                        f"{index}"
                    ),
                )

        st.info(
            "Tip: Only include keywords that genuinely "
            "match your skills or experience."
        )


    # ======================================================
    # ATS CHECKLIST
    # ======================================================

    with tool_tabs[5]:

        ats_checklist()


    # ======================================================
    # ATS RESUME TEMPLATES
    # ======================================================

    with tool_tabs[6]:

        ats_resume_templates()